import io
import csv
import json
import os
import re
import sqlite3
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock, patch

from cryptography.fernet import Fernet
from google.genai import types
from openpyxl import load_workbook

from report_export import build_csv, build_pdf, build_xlsx, DISCLAIMER
from commerce_signals import detect_commerce_signals, commerce_metadata


class ReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = TemporaryDirectory()
        cls.environment = patch.dict(os.environ, {
            "ENCRYPTION_KEY": Fernet.generate_key().decode(), "GEMINI_API_KEY": "",
        })
        cls.environment.start()
        import database
        cls.database = database
        cls.db_path = patch.object(database, "DB_PATH", Path(cls.directory.name) / "test.db")
        cls.db_path.start()
        import app
        cls.module = app

    @classmethod
    def tearDownClass(cls):
        cls.db_path.stop()
        cls.environment.stop()
        cls.directory.cleanup()

    def setUp(self):
        self.company = {
            "name": "Компания", "revenue": 1000, "tax_percent": 0.5,
            "employees": 0, "social_media": "высокая", "region": None, "sector": None,
        }
        self.data = self.module.normalize_report_input(
            self.module.analyze_company(self.company), self.company
        )
        self.report = {
            "company_name": "Компания", "revenue": 1000,
            "employees": 0, "social_media": "высокая", "score": 77,
            "recommendation": {"needs_inspection": True, "reason": "Нужна сверка показателей."},
        }

    def generate(self, text, data=None):
        client = Mock()
        client.models.generate_content.return_value = SimpleNamespace(text=text)
        with patch.object(self.module, "client", client):
            result = self.module.generate_ai_report(data or self.data)
        config = client.models.generate_content.call_args.kwargs["config"]
        self.assertEqual(types.GenerateContentConfig(**config).response_mime_type, "application/json")
        return result

    def test_structured_report_and_order(self):
        result = self.generate(json.dumps(self.report))
        self.assertEqual(result, self.report)
        self.assertEqual(list(result), list(self.report))

    def test_facts_and_score_cannot_be_invented(self):
        invented = {**self.report, "company_name": "Выдумано", "revenue": 999, "score": 99}
        self.assertEqual(self.generate(json.dumps(invented)), self.report)
        aggregate = self.module.normalize_report_input({"risk_score": 0, "rows": 10})
        result = self.generate(json.dumps(invented), aggregate)
        self.assertIsNone(result["company_name"])
        self.assertIsNone(result["revenue"])
        self.assertEqual(result["score"], 0)

    def test_fenced_json(self):
        self.assertEqual(self.generate("```json\n" + json.dumps(self.report) + "\n```"), self.report)

    def test_invalid_json_and_structure_fall_back(self):
        for text in ("<script>текст</script>", "null", "[]", "{}", "```",
                     json.dumps({**self.report, "score": 101}),
                     json.dumps({**self.report, "recommendation": {"needs_inspection": "false", "reason": "Нет"}})):
            with self.subTest(text=text):
                self.assertEqual(self.generate(text), text)

    def test_missing_key_empty_response_and_api_failure(self):
        with patch.object(self.module, "client", None):
            self.assertIn("не настроен", self.module.generate_ai_report(self.data))
        self.assertIn("пустой ответ", self.generate(""))
        client = Mock()
        client.models.generate_content.side_effect = RuntimeError("Секретные сведения")
        with patch.object(self.module, "client", client):
            result = self.module.generate_ai_report(self.data)
        self.assertIn("Не удалось", result)
        self.assertNotIn("Секретные сведения", result)

    def test_manual_route_persists_recommendation(self):
        client = Mock()
        client.models.generate_content.return_value = SimpleNamespace(text=json.dumps(self.report))
        with patch.object(self.module, "client", client):
            response = self.module.app.test_client().post("/api/manual-entry", json=self.company)
        self.assertEqual(response.status_code, 200)
        result = response.json
        self.assertEqual(result["ai_report"], self.report)
        saved = self.database.get_company(result["company"]["id"])
        self.assertEqual(json.loads(saved["recommendation"]), self.report["recommendation"])
        with sqlite3.connect(self.database.DB_PATH) as connection:
            name = connection.execute("SELECT name FROM companies WHERE id = ?", (saved["id"],)).fetchone()[0]
        self.assertNotEqual(name, self.company["name"])

    def test_uploads_use_normalized_input(self):
        for csv, expected in (
            ("name,revenue,tax_percent,employees,social_media\nКомпания,1000,0.5,0,высокая\n", "Компания"),
            ("name,revenue,tax_percent,employees\nПервая,1000,5,1\nВторая,2000,5,2\n", "Вторая"),
        ):
            with self.subTest(csv=csv), patch.object(self.module, "generate_ai_report", return_value=self.report) as generate:
                response = self.module.app.test_client().post("/api/analyze", data={
                    "file": (io.BytesIO(csv.encode()), "sample.csv"),
                })
                self.assertEqual(response.status_code, 200)
                normalized = generate.call_args.args[0]
                self.assertEqual(set(normalized), {"company", "analysis"})
                self.assertEqual(normalized["company"].get("name"), expected)

    def test_export_builders(self):
        csv_bytes = build_csv(self.report)
        self.assertTrue(csv_bytes.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(csv_bytes.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(len(rows), 8)
        self.assertEqual([row[0] for row in rows[:6]], [
            "Название компании", "Доход", "Количество сотрудников",
            "Активность в социальных сетях", "Оценка риска", "Рекомендация",
        ])
        self.assertEqual(rows[0][1], "Компания")
        self.assertEqual(rows[2][1], "0")
        self.assertIn("Проверка нужна: Да", rows[5][1])
        self.assertEqual(rows[6][1], DISCLAIMER)
        self.assertIn("UTC", rows[7][1])
        workbook = load_workbook(io.BytesIO(build_xlsx(self.report)))
        self.assertEqual(len(workbook.worksheets), 1)
        self.assertEqual(workbook.active["B1"].value, "Компания")
        self.assertEqual(workbook.active["B2"].value, 1000)
        self.assertEqual(workbook.active["B3"].value, 0)
        workbook.close()
        long_report = {**self.report, "recommendation": {
            "needs_inspection": False, "reason": "Текст проверки. " * 1000,
        }}
        for report in (self.report, long_report):
            pdf = build_pdf(report)
            self.assertTrue(pdf.startswith(b"%PDF-"))
            self.assertEqual(len(re.findall(rb"/Type\s*/Page\b", pdf)), 1)
            self.assertIn(b"/FontFile2", pdf)

    def test_commerce_keywords_and_threshold(self):
        self.assertEqual(detect_commerce_signals("ЦЕНА, заказ, ЗАКАЗ; доставка; в наличии"),
                         ["цена", "заказ", "доставка", "в наличии"])
        self.assertEqual(detect_commerce_signals("высокая"), [])
        for revenue in (0, 999, 1000):
            self.assertTrue(commerce_metadata("Доставка", revenue)["commerce_signal"])
        for missing in (None, "", float("nan"), float("inf")):
            self.assertFalse(commerce_metadata("Доставка", missing)["commerce_signal"])
        self.assertFalse(commerce_metadata("Доставка", 1001)["commerce_signal"])
        self.assertFalse(commerce_metadata("", 0)["commerce_signal"])

    def test_manual_commerce_persistence_and_report(self):
        company = {**self.company, "social_media": "Скидка, доставка, заказ"}
        with patch.object(self.module, "client", None):
            response = self.module.app.test_client().post("/api/manual-entry", json=company)
        self.assertEqual(response.status_code, 200)
        result = response.json
        self.assertTrue(result["company"]["commerce_signal"])
        self.assertEqual(result["company"]["commerce_keywords"], ["заказ", "доставка", "скидка"])
        self.assertEqual(result["risk_score"], 100)
        self.assertEqual(result["analysis"]["raw_score"], 65)
        self.assertIn("Обнаружены признаки торговой активности", result["ai_report"]["recommendation"]["reason"])
        analysis = self.module.analyze_company(company)
        data = self.module.normalize_report_input(analysis, company)
        report = self.generate(json.dumps(self.report), data)
        self.assertTrue(report["recommendation"]["needs_inspection"])
        self.assertIn("заказ, доставка, скидка", report["recommendation"]["reason"])
        self.assertEqual(list(report), list(self.report))

    def test_upload_commerce_matches_each_rows_revenue(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "commerce.csv"
            path.write_text("social_media,revenue,tax_percent,employees\nДоставка,1001,5,1\nЗаказ,1000,5,1\n", encoding="utf-8")
            result = self.module.analyze_data(path.as_posix())
            self.assertIsNone(result["risk_score"])
            self.assertFalse(result["records"][0]["analysis"]["commerce_signal"])
            self.assertEqual(result["records"][1]["analysis"]["risk_score"], 23)
            path.write_text("social_media,revenue\nДоставка,\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Строка 2"):
                self.module.analyze_data(path.as_posix())

    def test_phone_migration(self):
        with TemporaryDirectory() as directory, patch.object(self.database, "DB_PATH", Path(directory) / "legacy.db"):
            with sqlite3.connect(self.database.DB_PATH) as connection:
                connection.execute("CREATE TABLE companies (id INTEGER PRIMARY KEY, name TEXT)")
                connection.execute("INSERT INTO companies (id, name) VALUES (1, 'сохранено')")
            self.database.init_db()
            self.database.init_db()
            with sqlite3.connect(self.database.DB_PATH) as connection:
                self.assertEqual(connection.execute("SELECT name, phone FROM companies").fetchone(), ("сохранено", None))
                self.assertEqual([r[1] for r in connection.execute("PRAGMA table_info(companies)")].count("phone"), 1)

    def test_same_company_manual_csv_and_excel(self):
        from openpyxl import Workbook
        company = {**self.company, "phone": "+998001234567", "region": "Tashkent"}
        csv_text = "name,revenue,tax_paid,employees,social_media,phone,region\nКомпания,1000,5,0,высокая,+998001234567,Tashkent\n"
        workbook = Workbook()
        for row in csv.reader(io.StringIO(csv_text)):
            workbook.active.append(row)
        excel = io.BytesIO()
        workbook.save(excel)
        client = self.module.app.test_client()
        with patch.object(self.module, "client", None):
            manual = client.post("/api/manual-entry", json=company).json
            for extension, payload in (("csv", csv_text.encode()), ("xlsx", excel.getvalue())):
                response = client.post("/api/analyze", data={"file": (io.BytesIO(payload), "input." + extension)})
                self.assertEqual(response.status_code, 200, response.json)
                result = response.json
                self.assertEqual(result["analysis"], manual["analysis"])
                self.assertEqual(result["analysis"]["risk_score"], 77)
                self.assertEqual(result["analysis"]["risk_tier"], "high")
                saved = self.database.get_company(result["company"]["id"])
                self.assertEqual(saved["phone"], company["phone"])
                self.assertEqual(saved["tax_percent"], 0.5)
                self.assertEqual(saved["scoring_version"], self.module.SCORING_VERSION)

    def test_batch_quality_does_not_change_individual_scores(self):
        content = "name,revenue,tax_paid,employees,region\nA,1000,5,0,\nA,1000,5,0,\nB,1000000,50000,100,Tashkent\n"
        before = len(self.database.get_all_companies())
        with patch.object(self.module, "client", None):
            response = self.module.app.test_client().post("/api/analyze", data={"file": (io.BytesIO(content.encode()), "batch.csv")})
        self.assertEqual(response.status_code, 200)
        result = response.json
        self.assertIsNone(result["analysis"]["risk_score"])
        self.assertIsNone(result["company"])
        self.assertEqual(result["dataset_analysis"]["duplicates"], 1)
        self.assertEqual(result["dataset_analysis"]["missing_values"], 2)
        self.assertEqual([item["analysis"]["risk_score"] for item in result["results"]], [77, 77, 0])
        self.assertEqual(len(self.database.get_all_companies()), before + 3)
        for item in result["results"]:
            saved = self.database.get_company(item["company"]["id"])
            self.assertEqual(saved["risk_score"], item["analysis"]["risk_score"])

    def test_invalid_batch_never_saves_partial_records(self):
        before = len(self.database.get_all_companies())
        client = self.module.app.test_client()
        for row in ("B,inf,5,0", "B,-1,5,0", "B,1000,5,0.5", "B,1000,,0", "B,,5,0"):
            content = "name,revenue,tax_percent,employees\nA,1000,0.5,0\n" + row + "\n"
            with self.subTest(row=row), patch.object(self.module, "generate_ai_report") as generate:
                response = client.post("/api/analyze", data={"file": (io.BytesIO(content.encode()), "invalid.csv")})
                self.assertEqual(response.status_code, 400)
                self.assertIn("Строка 3", response.json["error"])
                generate.assert_not_called()
                self.assertEqual(len(self.database.get_all_companies()), before)

    def test_scoring_boundaries_and_missing_revenue(self):
        from scoring import analyze_company, normalize_company, risk_tier
        for tax, raw in ((0.99, 30), (1, 15), (2.99, 15), (3, 0)):
            self.assertEqual(analyze_company({**self.company, "employees": 1, "tax_percent": tax})["raw_score"], raw)
        for score, tier in ((0, "low"), (29, "low"), (30, "medium"), (70, "medium"), (71, "high"), (100, "high")):
            self.assertEqual(risk_tier(score), tier)
        maximum = analyze_company({**self.company, "social_media": "доставка"})
        self.assertEqual((maximum["raw_score"], maximum["risk_score"]), (65, 100))
        zero = analyze_company({**self.company, "revenue": 0, "tax_percent": None, "employees": 2, "social_media": "ВЫСОКАЯ"})
        self.assertEqual(zero["raw_score"], 32)
        with self.assertRaises(ValueError):
            normalize_company({**self.company, "tax_paid": 500})
        with self.assertRaises(ValueError):
            analyze_company({**self.company, "revenue": None})

    def test_legacy_scores_excluded_from_summary(self):
        companies = [
            {"region": "Tashkent", "risk_score": 100, "scoring_version": None},
            {"region": "Tashkent", "risk_score": 0, "scoring_version": self.module.SCORING_VERSION},
        ]
        with patch.object(self.module, "get_all_companies", return_value=companies):
            region = self.module.app.test_client().get("/api/regions/summary").json["regions"][0]
        self.assertEqual(region["company_count"], 1)
        self.assertEqual(region["average_risk_score"], 0)

    def test_sample_file_and_home_page(self):
        result = self.module.analyze_data(str(Path(__file__).resolve().parents[1] / "sample_data.csv"))
        self.assertEqual(len(result["records"]), 11)
        for record in result["records"]:
            self.assertEqual(record["analysis"]["risk_score"], 0)
            self.assertIsNotNone(record["company"]["tax_percent"])
        response = self.module.app.test_client().get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn('id="companyResultSelect"', response.get_data(as_text=True))
        self.assertIn('window.riskBands =', response.get_data(as_text=True))

    def test_batch_insert_rolls_back_on_failure(self):
        before = len(self.database.get_all_companies())
        with self.assertRaises(ValueError):
            self.database.insert_companies([
                {**self.company, "source": "upload"},
                {**self.company, "source": "invalid"},
            ])
        self.assertEqual(len(self.database.get_all_companies()), before)

    def test_encrypted_links_and_graph(self):
        with TemporaryDirectory() as directory, patch.object(self.database, "DB_PATH", Path(directory) / "links.db"):
            self.database.init_db()
            first = self.database.insert_company({"name": "Первая", "phone": " +99890 ", "social_media": " Профиль ", "source": "manual"})
            second = self.database.insert_company({"name": "Вторая", "phone": "+99890", "social_media": "профиль", "source": "manual"})
            self.database.insert_company({"name": "Без связи", "phone": " ", "social_media": "", "source": "manual"})
            self.database.insert_company({"source": "manual"})
            self.assertEqual(self.database.get_company(first)["phone"], " +99890 ")
            with sqlite3.connect(self.database.DB_PATH) as connection:
                encrypted = connection.execute("SELECT phone FROM companies WHERE id = ?", (first,)).fetchone()[0]
                self.assertNotIn("+99890", encrypted)
            pairs = self.database.find_linked_companies()
            self.assertEqual(len(pairs), 2)
            self.assertEqual({p["reason"] for p in pairs}, {"same phone", "same social media"})
            result = self.module.app.test_client().get("/api/graph")
            self.assertEqual(result.status_code, 200)
            self.assertEqual({n["id"] for n in result.json["nodes"]}, {first, second})
            self.assertEqual(len(result.json["edges"]), 2)
            self.assertNotIn("phone", result.json["nodes"][0])
            self.assertNotIn("+99890", result.get_data(as_text=True))

    def test_manual_phone_and_empty_graph(self):
        client = self.module.app.test_client()
        with patch.object(self.module, "client", None):
            result = client.post("/api/manual-entry", json={**self.company, "phone": " +998 90 "})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json["company"]["phone"], "+998 90")
        self.assertEqual(client.post("/api/manual-entry", json={**self.company, "phone": 123}).status_code, 400)
        with patch.object(self.module, "find_linked_companies", return_value=[]):
            self.assertEqual(client.get("/api/graph").json, {"nodes": [], "edges": [], "risk_bands": self.module.RISK_BANDS})
        with patch.object(self.module, "find_linked_companies", side_effect=ValueError("Секрет")):
            self.assertEqual(client.get("/api/graph").status_code, 500)

    def test_regions_summary(self):
        companies = [
            {"region": " Ташкент ", "risk_score": 70},
            {"region": "ташкент", "risk_score": 71},
            {"region": "Ташкент", "risk_score": None},
            {"region": "Самарканд", "risk_score": 0},
            {"region": None, "risk_score": 100},
            {"region": "  ", "risk_score": None},
            {"region": "Бухара", "risk_score": None},
        ]
        companies = [{**item, "scoring_version": self.module.SCORING_VERSION} for item in companies]
        with patch.object(self.module, "get_all_companies", return_value=companies):
            response = self.module.app.test_client().get("/api/regions/summary")
        self.assertEqual(response.status_code, 200)
        regions = {row["region"]: row for row in response.json["regions"]}
        self.assertEqual(regions["Ташкент"], {
            "region": "Ташкент", "company_count": 3,
            "average_risk_score": 70.5, "high_risk_count": 1,
        })
        self.assertEqual(regions["Самарканд"]["average_risk_score"], 0)
        self.assertEqual(regions["Регион не указан"]["company_count"], 2)
        self.assertIsNone(regions["Бухара"]["average_risk_score"])
        with patch.object(self.module, "get_all_companies", return_value=[]):
            self.assertEqual(self.module.app.test_client().get("/api/regions/summary").json["regions"], [])
        with patch.object(self.module, "get_all_companies", side_effect=ValueError("Ошибка ключа")):
            self.assertEqual(self.module.app.test_client().get("/api/regions/summary").status_code, 500)

    def test_export_formula_protection(self):
        record = {**self.report, "company_name": '=HYPERLINK("https://example.com")'}
        csv_rows = list(csv.reader(io.StringIO(build_csv(record).decode("utf-8-sig")), delimiter=";"))
        self.assertTrue(csv_rows[0][1].startswith("'="))
        workbook = load_workbook(io.BytesIO(build_xlsx(record)))
        self.assertEqual(workbook.active["B1"].data_type, "s")
        self.assertEqual(workbook.active["B1"].value, record["company_name"])
        workbook.close()

    def test_download_for_both_flows(self):
        client = self.module.app.test_client()
        with patch.object(self.module, "generate_ai_report", return_value=self.report):
            manual = client.post("/api/manual-entry", json=self.company)
            uploaded = client.post("/api/analyze", data={"file": (
                io.BytesIO("name,revenue,tax_percent,employees,social_media\nКомпания,1000,0.5,0,высокая\n".encode()), "sample.csv",
            )})
        for result in (manual, uploaded):
            self.assertEqual(result.status_code, 200)
            company_id = result.json["company"]["id"]
            for extension, mime in (("pdf", "application/pdf"), ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), ("csv", "text/csv")):
                response = client.get(f"/api/report/{company_id}/download?format={extension}")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.mimetype, mime)
                self.assertIn(f"attachment; filename=otchet_{company_id}.{extension}", response.headers["Content-Disposition"])
                self.assertEqual(response.headers["Cache-Control"], "no-store")
                if extension == "csv":
                    self.assertIn("Компания", response.data.decode("utf-8-sig"))
                    self.assertNotIn("gAAAA", response.data.decode("utf-8-sig"))

    def test_download_errors_and_unavailable_recommendation(self):
        client = self.module.app.test_client()
        self.assertEqual(client.get("/api/report/999999/download?format=csv").status_code, 404)
        self.assertEqual(client.get("/api/report/1/download?format=exe").status_code, 400)
        with patch.object(self.module, "client", None):
            result = client.post("/api/manual-entry", json=self.company)
        company_id = result.json["company"]["id"]
        response = client.get(f"/api/report/{company_id}/download?format=csv")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Проверка нужна: Не определено", response.data.decode("utf-8-sig"))


if __name__ == "__main__":
    unittest.main()
