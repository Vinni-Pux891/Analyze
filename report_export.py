import csv
import math
import re
from datetime import datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph


DISCLAIMER = "Risk Score — аналитический индикатор, не юридическое обвинение"
FONT_PATH = Path(__file__).resolve().parent / "static/fonts/Montserrat-Regular.ttf"
pdfmetrics.registerFont(TTFont("ReportFont", str(FONT_PATH)))


def _text(value):
    if value is None or value == "":
        return "Не указано"
    if isinstance(value, float) and not math.isfinite(value):
        return "Не указано"
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))


def _rows(record):
    recommendation = record.get("recommendation") or {}
    decision = recommendation.get("needs_inspection")
    answer = "Да" if decision is True else "Нет" if decision is False else "Не определено"
    reason = _text(recommendation.get("reason"))
    score = record.get("score")
    return [
        ("Название компании", record.get("company_name")),
        ("Доход", record.get("revenue")),
        ("Количество сотрудников", record.get("employees")),
        ("Активность в социальных сетях", record.get("social_media")),
        ("Оценка риска", f"{score}/100" if score is not None else None),
        ("Рекомендация", f"Проверка нужна: {answer}. {reason}"),
        ("Примечание", DISCLAIMER + (f". Методика: {record['scoring_version']}" if record.get("scoring_version") else "")),
        ("Дата формирования", datetime.now(timezone.utc).strftime("%d.%m.%Y %H:%M:%S (UTC)")),
    ]


def build_pdf(record: dict) -> bytes:
    output = BytesIO()
    page = canvas.Canvas(output, pagesize=A4)
    page.setTitle("Результат анализа")
    width, height = A4
    margin = 40
    style = ParagraphStyle("Отчёт", fontName="ReportFont", fontSize=11, leading=16)
    paragraphs = []
    for label, value in _rows(record):
        text = escape(f"{label}: {_text(value)}").replace("\n", "<br/>")
        paragraph = Paragraph(text, style)
        _, paragraph_height = paragraph.wrap(width - 2 * margin, height)
        paragraphs.append((paragraph, paragraph_height))
    total_height = sum(size + 14 for _, size in paragraphs)
    # Масштабирование сохраняет весь текст на одной странице без обрезания.
    scale = min(1, (height - 2 * margin) / total_height)
    page.translate(margin, height - margin)
    page.scale(scale, scale)
    y = 0
    for paragraph, size in paragraphs:
        y -= size
        paragraph.drawOn(page, 0, y)
        y -= 14
    page.showPage()
    page.save()
    return output.getvalue()


def build_xlsx(record: dict) -> bytes:
    output = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Результат анализа"
    for row_number, (label, value) in enumerate(_rows(record), 1):
        sheet.cell(row_number, 1, label).font = Font(bold=True)
        cell = sheet.cell(row_number, 2)
        if type(value) in (int, float) and math.isfinite(value):
            cell.value = value
        else:
            text = _text(value)
            if len(text) > 32767:
                raise ValueError("Текст слишком длинный для ячейки Excel. Выберите PDF или CSV.")
            cell.value = text
            # Пользовательские значения сохраняются как текст, а не формулы.
            cell.data_type = "s"
        for item in sheet[row_number]:
            item.alignment = Alignment(wrap_text=True, vertical="top")
        lines = sum(max(1, math.ceil(len(line) / 65)) for line in _text(value).split("\n"))
        sheet.row_dimensions[row_number].height = min(409, max(32, 16 * lines))
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 85
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = "A1:B8"
    workbook.save(output)
    return output.getvalue()


def build_csv(record: dict) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";")
    for label, value in _rows(record):
        text = _text(value)
        # Предотвращаем выполнение формул при открытии CSV в Excel.
        if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
            text = "'" + text
        writer.writerow((label, text))
    return output.getvalue().encode("utf-8-sig")
